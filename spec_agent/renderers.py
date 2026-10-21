"""Render a ``TechSpec`` to Markdown, Word (.docx) and JSON.

Both document formats are produced from one intermediate list of blocks, so the Markdown and the
Word export always contain the same sections in the same order.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import BinaryIO, Literal

from .linter import LintReport
from .llm import OFFLINE_BANNER
from .schema import ScoreReport, TechSpec

BlockKind = Literal["title", "h2", "h3", "p", "bullets", "table", "callout"]


@dataclass
class Block:
    kind: BlockKind
    text: str = ""
    items: list[str] | None = None
    headers: list[str] | None = None
    rows: list[list[str]] | None = None


def _yn(flag: bool) -> str:
    return "yes" if flag else "no"


def build_blocks(spec: TechSpec, score: ScoreReport | None = None,
                 lint: LintReport | None = None, guardrail_notes: list[str] | None = None
                 ) -> list[Block]:
    B = Block
    out: list[Block] = [B("title", f"Technical Spec: {spec.title}")]
    if spec.meta.mode == "offline-demo":
        out.append(B("callout", f"**Offline demo output.** {OFFLINE_BANNER}. Requirements are "
                                "restated from the request; everything labelled *assumption* "
                                "and every open question needs human review."))
    meta_rows = [
        ["Status", f"Draft, generated {spec.meta.generated_at}"],
        ["Mode", f"{spec.meta.mode} (model: {spec.meta.model})"],
        ["Pipeline", "extract → draft → critique → revise" if spec.meta.pipeline == "full"
         else "single pass (extract → draft, no critique)"],
    ]
    if spec.meta.source_file:
        meta_rows.append(["Source", spec.meta.source_file])
    if score:
        meta_rows.append(["Quality score", f"{score.total:g} / {score.max_total:g}"])
    meta_rows.append(["Open questions", f"{len(spec.open_questions)} "
                      f"({sum(q.blocking for q in spec.open_questions)} blocking)"])
    out.append(B("table", headers=["Field", "Value"], rows=meta_rows))

    out += [B("h2", "1. Overview"), B("p", spec.overview),
            B("h2", "2. Problem statement"), B("p", spec.problem_statement),
            B("h2", "3. Goals"), B("bullets", items=spec.goals or ["None stated."]),
            B("h2", "4. Non-goals"), B("bullets", items=spec.non_goals or ["None stated."])]

    out.append(B("h2", "5. User stories"))
    for s in spec.user_stories:
        out.append(B("p", f"**{s.id}** ({s.priority}): As a {s.as_a}, I want {s.i_want}, so "
                          f"that {s.so_that}."))
        out.append(B("bullets", items=[f"**Given** {ac.given} **when** {ac.when} **then** "
                                       f"{ac.then}" for ac in s.acceptance_criteria]
                     or ["No acceptance criteria."]))

    out += [B("h2", "6. Requirements"), B("h3", "6.1 Functional"),
            B("table", headers=["ID", "Requirement", "Priority", "Source", "Refs"],
              rows=[[r.id, r.description, r.priority, r.source, ", ".join(r.refs)]
                    for r in spec.functional_requirements]),
            B("h3", "6.2 Non-functional"),
            B("table", headers=["ID", "Category", "Requirement", "Priority", "Source", "Refs"],
              rows=[[r.id, r.category, r.description, r.priority, r.source, ", ".join(r.refs)]
                    for r in spec.non_functional_requirements])]

    a = spec.analytics
    out += [B("h2", "7. Analytics & tracking plan"),
            B("p", f"Naming convention: {a.naming_convention}."),
            B("h3", "7.1 Global properties (sent with every event)"),
            B("table", headers=["Property", "Type", "Required", "Description", "Allowed values"],
              rows=[[f"`{p.name}`", p.type.value, _yn(p.required), p.description,
                     ", ".join(p.allowed_values or [])] for p in a.global_properties]),
            B("h3", "7.2 Events"),
            B("table", headers=["Event", "Status", "Trigger", "Platforms", "Requirements"],
              rows=[[f"`{e.name}`", e.status, e.trigger, ", ".join(e.platforms),
                     ", ".join(e.related_requirements)] for e in a.events]),
            B("h3", "7.3 Event properties")]
    for e in a.events:
        out.append(B("p", f"**`{e.name}`** ({e.status}): {e.description}"))
        out.append(B("table", headers=["Property", "Type", "Required", "Description",
                                       "Allowed values"],
                     rows=[[f"`{p.name}`", p.type.value, _yn(p.required), p.description,
                            ", ".join(p.allowed_values or [])] for p in e.properties]))
    out += [B("h3", "7.4 Success metrics"),
            B("table", headers=["Metric", "Definition", "Target", "Events"],
              rows=[[m.name, m.definition, m.target, ", ".join(f"`{x}`" for x in m.events)]
                    for m in a.success_metrics])]

    out.append(B("h2", "8. Data model & API changes"))
    if spec.data_model_changes:
        out.append(B("table", headers=["Entity", "Change", "Fields", "Description", "Source"],
                     rows=[[d.entity, d.change_type, ", ".join(f"`{f}`" for f in d.fields),
                            d.description, d.source] for d in spec.data_model_changes]))
    if spec.api_changes:
        out.append(B("table", headers=["Method", "Path", "Change", "Description", "Source"],
                     rows=[[c.method, f"`{c.path}`", c.change_type, c.description, c.source]
                           for c in spec.api_changes]))
    if not (spec.data_model_changes or spec.api_changes):
        out.append(B("p", "None identified."))

    out += [B("h2", "9. Dependencies"),
            B("table", headers=["Dependency", "Type", "Description", "Source"],
              rows=[[d.name, d.type, d.description, d.source] for d in spec.dependencies]),
            B("h2", "10. Edge cases"),
            B("table", headers=["ID", "Scenario", "Expected behaviour", "Source"],
              rows=[[e.id, e.scenario, e.expected_behavior, e.source] for e in spec.edge_cases])]

    out.append(B("h2", "11. Rollout & experiment plan"))
    if spec.rollout:
        r = spec.rollout
        out.append(B("p", f"Feature flag: `{r.feature_flag}` (source: {r.source})."))
        out.append(B("table", headers=["Phase", "Audience", "Exposure", "Exit criteria"],
                     rows=[[p.name, p.audience, f"{p.exposure_percent}%", p.exit_criteria]
                           for p in r.phases]))
        if r.experiment:
            x = r.experiment
            out.append(B("bullets", items=[
                f"**Hypothesis:** {x.hypothesis}", f"**Variants:** {'; '.join(x.variants)}",
                f"**Randomisation unit:** {x.randomization_unit}",
                f"**Primary metric:** {x.primary_metric}",
                f"**Guardrails:** {', '.join(x.guardrail_metrics) or 'none'}",
                f"**Minimum detectable effect:** {x.minimum_detectable_effect}"]))
        out.append(B("p", f"Rollback: {r.rollback_plan}"))
    else:
        out.append(B("p", "No rollout plan."))

    out += [B("h2", "12. Test plan"),
            B("table", headers=["ID", "Type", "Title", "Covers", "Events verified", "Expected"],
              rows=[[t.id, t.type, t.title, ", ".join(t.covers),
                     ", ".join(f"`{e}`" for e in t.events_verified), t.expected_result]
                    for t in spec.test_plan]),
            B("h2", "13. Risks"),
            B("table", headers=["Risk", "Likelihood", "Impact", "Mitigation"],
              rows=[[r.description, r.likelihood, r.impact, r.mitigation] for r in spec.risks]),
            B("h2", "14. Assumptions"),
            B("table", headers=["ID", "Assumption", "Validate with"],
              rows=[[x.id, x.statement, x.validate_with] for x in spec.assumptions]),
            B("h2", "15. Open questions"),
            B("table", headers=["ID", "Question", "Why it matters", "Owner", "Blocking"],
              rows=[[q.id, q.question, q.why_it_matters, q.suggested_owner, _yn(q.blocking)]
                    for q in spec.open_questions])]

    if score:
        out += [B("h2", "Appendix A. Quality score"),
                B("p", f"Deterministic rubric score: **{score.total:g} / {score.max_total:g}**."),
                B("table", headers=["Category", "Score", "Max"],
                  rows=[[c.name, f"{c.score:g}", f"{c.max_score:g}"] for c in score.categories])]
        items = [f"**{f.severity}** ({f.location}): {f.message}" for f in score.findings]
        out.append(B("bullets", items=items or ["No findings."]))
    if lint is not None:
        out += [B("h2", "Appendix B. Event dictionary lint"),
                B("p", f"Checked {lint.checked_events} events against {lint.dictionary_size} "
                       "dictionary entries.")]
        out.append(B("table", headers=["Severity", "Event", "Rule", "Finding", "Suggestion"],
                     rows=[[f.severity, f"`{f.event}`", f.rule, f.message, f.suggestion]
                           for f in lint.findings]) if lint.findings
                   else B("p", "No findings."))
    if guardrail_notes:
        out += [B("h2", "Appendix C. Provenance guardrail"), B("bullets", items=guardrail_notes)]
    return out


# --- Markdown ----------------------------------------------------------------------------------


def _cell(text: str) -> str:
    return text.replace("|", "\\|").replace("\n", " ")


def to_markdown(blocks: list[Block]) -> str:
    lines: list[str] = []
    for b in blocks:
        if b.kind == "title":
            lines += [f"# {b.text}", ""]
        elif b.kind == "h2":
            lines += [f"## {b.text}", ""]
        elif b.kind == "h3":
            lines += [f"### {b.text}", ""]
        elif b.kind == "p":
            lines += [b.text, ""]
        elif b.kind == "callout":
            lines += [f"> {b.text}", ""]
        elif b.kind == "bullets":
            lines += [f"- {i}" for i in b.items or []] + [""]
        elif b.kind == "table":
            if not b.rows:
                lines += ["_None._", ""]
                continue
            lines.append("| " + " | ".join(b.headers or []) + " |")
            lines.append("|" + "---|" * len(b.headers or []))
            lines += ["| " + " | ".join(_cell(c) for c in row) + " |" for row in b.rows]
            lines.append("")
    return "\n".join(lines).rstrip() + "\n"


# --- Word --------------------------------------------------------------------------------------

_INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+\*)")


def _add_runs(paragraph, text: str) -> None:  # type: ignore[no-untyped-def]
    """Very small inline-markdown renderer: **bold**, *italic*, `code`."""
    from docx.shared import Pt

    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def _column_widths(headers: list[str], rows: list[list[str]], total: float = 7.1
                   ) -> list[float]:
    """Share the page width by content length (capped) so long text columns get the room."""
    weights = [min(60, max(len(h), *(len(r[i]) for r in rows))) + 6
               for i, h in enumerate(headers)]
    return [total * w / sum(weights) for w in weights]


def to_docx(blocks: list[Block], path: str | Path | BinaryIO) -> str | Path | BinaryIO:
    """Write a .docx to a path or a binary stream (e.g. ``io.BytesIO`` for downloads)."""
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    for section in doc.sections:
        section.left_margin = section.right_margin = Inches(0.7)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    for b in blocks:
        if b.kind == "title":
            doc.add_heading(b.text, level=0)
        elif b.kind in ("h2", "h3"):
            doc.add_heading(b.text, level=1 if b.kind == "h2" else 2)
        elif b.kind == "p":
            _add_runs(doc.add_paragraph(), b.text)
        elif b.kind == "callout":
            para = doc.add_paragraph()
            _add_runs(para, b.text)
            for run in para.runs:
                run.font.color.rgb = RGBColor(0x9A, 0x4D, 0x00)
        elif b.kind == "bullets":
            for item in b.items or []:
                _add_runs(doc.add_paragraph(style="List Bullet"), item)
        elif b.kind == "table":
            if not b.rows:
                doc.add_paragraph("None.")
                continue
            table = doc.add_table(rows=1, cols=len(b.headers or []))
            table.style = "Light Grid Accent 1"
            table.autofit = False
            for cell, header in zip(table.rows[0].cells, b.headers or []):
                cell.text = ""
                cell.paragraphs[0].add_run(header).bold = True
            for row in b.rows:
                cells = table.add_row().cells
                for cell, value in zip(cells, row):
                    cell.text = ""
                    _add_runs(cell.paragraphs[0], value)
            widths = _column_widths(b.headers or [], b.rows)
            # Set both the grid (used by LibreOffice/pandoc) and per-cell widths (used by Word).
            for grid_col, width in zip(table._tbl.tblGrid.gridCol_lst, widths):
                grid_col.w = Inches(width)
            for col, width in zip(table.columns, widths):
                for cell in col.cells:
                    cell.width = Inches(width)
            doc.add_paragraph()
    if isinstance(path, (str, Path)):
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(path))
    else:
        doc.save(path)
    return path


# --- Facade ------------------------------------------------------------------------------------


def render_markdown(spec: TechSpec, score: ScoreReport | None = None,
                    lint: LintReport | None = None, guardrail_notes: list[str] | None = None
                    ) -> str:
    return to_markdown(build_blocks(spec, score, lint, guardrail_notes))


def render_docx(spec: TechSpec, path: str | Path | BinaryIO, score: ScoreReport | None = None,
                lint: LintReport | None = None, guardrail_notes: list[str] | None = None
                ) -> str | Path | BinaryIO:
    return to_docx(build_blocks(spec, score, lint, guardrail_notes), path)


def render_json(spec: TechSpec) -> str:
    return spec.model_dump_json(indent=2)


def write_outputs(spec: TechSpec, out_dir: str | Path, stem: str, formats: list[str],
                  score: ScoreReport | None = None, lint: LintReport | None = None,
                  guardrail_notes: list[str] | None = None) -> list[Path]:
    """Write the requested formats (``md``, ``docx``, ``json``) and return the paths written."""
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    written: list[Path] = []
    for fmt in formats:
        if fmt == "md":
            p = out / f"{stem}.md"
            p.write_text(render_markdown(spec, score, lint, guardrail_notes), encoding="utf-8")
        elif fmt == "docx":
            p = out / f"{stem}.docx"
            render_docx(spec, p, score, lint, guardrail_notes)
        elif fmt == "json":
            p = out / f"{stem}.json"
            p.write_text(render_json(spec), encoding="utf-8")
        else:
            raise ValueError(f"Unknown format '{fmt}' (use md, docx, json)")
        written.append(p)
    return written
