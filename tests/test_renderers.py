import io
import json

from docx import Document

from spec_agent.renderers import render_json, render_markdown, write_outputs
from spec_agent.schema import TechSpec
from spec_agent.scorer import score_spec


def test_markdown_contains_all_sections_and_offline_label(good_spec):
    md = render_markdown(good_spec, score_spec(good_spec))
    assert md.startswith("# Technical Spec: AutoPay enrollment from the bill page")
    assert "Offline demo output" in md
    for heading in ["## 1. Overview", "## 5. User stories", "### 6.1 Functional",
                    "## 7. Analytics & tracking plan", "### 7.2 Events", "## 11. Rollout",
                    "## 12. Test plan", "## 15. Open questions", "## Appendix A. Quality score"]:
        assert heading in md
    assert "| `autopay_enrollment_submitted` | new |" in md
    assert "**Given** a signed-in account owner" in md


def test_llm_mode_has_no_offline_label(good_spec):
    good_spec.meta.mode = "llm"
    assert "Offline demo output" not in render_markdown(good_spec)


def test_pipe_characters_are_escaped_in_tables(good_spec):
    good_spec.functional_requirements[0].description = "Show A | B options."
    assert "Show A \\| B options." in render_markdown(good_spec)


def test_write_outputs_produces_md_docx_json(good_spec, tmp_path):
    paths = write_outputs(good_spec, tmp_path, "autopay", ["md", "docx", "json"],
                          score_spec(good_spec))
    assert [p.name for p in paths] == ["autopay.md", "autopay.docx", "autopay.json"]
    assert all(p.stat().st_size > 0 for p in paths)
    doc = Document(str(tmp_path / "autopay.docx"))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert "7. Analytics & tracking plan" in headings and "12. Test plan" in headings
    assert len(doc.tables) >= 10
    restored = TechSpec.model_validate(json.loads((tmp_path / "autopay.json").read_text()))
    assert restored.title == good_spec.title


def test_docx_can_be_rendered_to_a_stream(good_spec):
    from spec_agent.renderers import render_docx

    buf = io.BytesIO()
    render_docx(good_spec, buf)
    assert buf.getvalue()[:2] == b"PK"  # .docx is a zip container


def test_render_json_is_valid(good_spec):
    assert json.loads(render_json(good_spec))["meta"]["mode"] == "offline-demo"
